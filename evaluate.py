"""
Step 4: Evaluate — เปรียบเทียบ Baseline vs Fine-tuned Typhoon OCR
==================================================================
วัดผลด้วย 4 metrics:
  1. CER  (Character Error Rate)     ← ระดับตัวอักษร
  2. WER  (Word Error Rate)          ← ระดับคำ
  3. Field Accuracy                  ← ระดับข้อมูลสำคัญ (เหมาะ Thesis)
  4. Exact Match Rate (EMR)          ← ตรงทั้งหมดหรือเปล่า

อ้างอิง:
  - CER/WER: Docusumo OCR Accuracy Guide (2025)
  - Field Accuracy: OCR Accuracy Benchmarks 2025 (Medium)
  - Typhoon OCR Benchmark: arXiv:2601.14722

ติดตั้ง:
  pip install transformers peft torch pillow tqdm jiwer
  pip install openpyxl   ← สำหรับ Export Excel

โครงสร้างโฟลเดอร์:
  dataset/formatted/test.jsonl   ← จาก format_dataset.py
  models/best_model/             ← จาก finetune.py
"""

import os
import json
import base64
import torch
import re
from io import BytesIO
from pathlib import Path
from dataclasses import dataclass, field
from typing import Optional
from datetime import datetime

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from PIL import Image, ImageOps
from tqdm import tqdm
from jiwer import cer, wer   # pip install jiwer

from transformers import BitsAndBytesConfig
from peft import PeftModel

import schema   # single source of truth ของ field/type/canonical JSON

# model/processor class เลือกอัตโนมัติตาม BASE_MODEL_ID (Qwen2.5-VL หรือ Qwen3-VL)
ModelClass     = schema.get_model_class()
ProcessorClass = schema.get_processor_class()


# ══════════════════════════════════════════════
# CONFIG
# ══════════════════════════════════════════════
@dataclass
class Config:
    # Paths
    TEST_JSONL:       str = "dataset/formatted/test.jsonl"
    # ต้องตรงกับ base ที่ adapter เทรนมา (schema = single source of truth)
    BASE_MODEL_ID:    str = schema.BASE_MODEL_ID
    FINETUNED_DIR:    str = "models/best_model"
    OUTPUT_DIR:       str = "results"

    # Generation
    MAX_NEW_TOKENS:   int = 768          # JSON ใบหลายรายการ 512 อาจไม่พอ
    # image resolution ต้องตรงกับตอน train (Finetune.py) ไม่งั้น LoRA เพี้ยน
    MIN_PIXELS:       int = 64 * 32 * 32
    MAX_PIXELS:       int = 320 * 32 * 32

    # Fields ที่วัด Field Accuracy — ดึงจาก schema.py (scalar + summary)
    # items วัดแยกต่างหาก (calc_items_metrics)
    FIELDS: list = field(default_factory=lambda: [
        {"key": k, "label": schema.LABELS.get(k, k), "type": schema.FIELD_TYPES.get(k, "fuzzy")}
        for k in (schema.SCALAR_FIELDS + schema.SUMMARY_FIELDS)
    ])


cfg = Config()


# ══════════════════════════════════════════════
# METRICS
# ══════════════════════════════════════════════

def calc_cer(prediction: str, ground_truth: str) -> float:
    """
    Character Error Rate (CER)

    สูตร: CER = (Substitutions + Deletions + Insertions) / Total Characters
    อ้างอิง: Levenshtein Distance บน character level

    ตัวอย่าง:
      GT:   "ใบกำกับภาษี"  (11 ตัวอักษร)
      Pred: "ใบกากับภาษี"  (ผิด 1 ตัว)
      CER:  1/11 = 0.091 = 9.1%

    เกณฑ์ (อ้างอิงงานวิจัย):
      CER < 2%  → ดีมาก (98%+ accuracy)
      CER < 5%  → ดี
      CER < 10% → พอใช้
      CER > 10% → ต้องปรับปรุง
    """
    if not ground_truth:
        return 0.0
    if not prediction:
        return 1.0
    return cer(ground_truth, prediction)


def calc_wer(prediction: str, ground_truth: str) -> float:
    """
    Word Error Rate (WER)

    สูตร: WER = (Substitutions + Deletions + Insertions) / Total Words
    WER มักสูงกว่า CER เสมอ (1 คำผิด = หลายตัวอักษร)

    เหมาะกับ: ข้อความทั่วไป, ชื่อบริษัท
    ไม่เหมาะกับ: ตัวเลข, รหัส (ใช้ CER หรือ Exact Match แทน)
    """
    if not ground_truth:
        return 0.0
    if not prediction:
        return 1.0
    return wer(ground_truth, prediction)


def calc_field_accuracy(pred_fields: dict, gt_fields: dict, field_cfg: dict) -> dict:
    """
    วัด Field-level Accuracy จาก fields ที่ parse จาก JSON แล้ว (ไม่ใช่ regex markdown)

    3 ประเภทการเปรียบเทียบ:
    1. exact   → ต้องตรงทุกตัวอักษร (เลขภาษี 13 หลัก, เลขที่ใบกำกับ)
    2. fuzzy   → ยอมรับความต่างเล็กน้อย (ชื่อ/ที่อยู่ — CER < 10%)
    3. numeric → ตัด comma แล้วเปรียบเทียบเชิงตัวเลข
    """
    key   = field_cfg["key"]
    ftype = field_cfg["type"]

    gt_val   = str(gt_fields.get(key, "")).strip()
    pred_val = str(pred_fields.get(key, "")).strip()

    if not gt_val:
        return {"matched": None, "gt": gt_val, "pred": pred_val}  # ข้ามถ้า GT ว่าง

    if ftype == "exact":
        matched = pred_val == gt_val

    elif ftype == "numeric":
        gt_num   = gt_val.replace(",", "").replace(" ", "")
        pred_num = pred_val.replace(",", "").replace(" ", "")
        try:
            matched = abs(float(pred_num) - float(gt_num)) < 0.01
        except ValueError:
            matched = False

    else:  # fuzzy
        gt_norm   = re.sub(r"\s+", " ", gt_val.lower()).strip()
        pred_norm = re.sub(r"\s+", " ", pred_val.lower()).strip()
        matched   = calc_cer(pred_norm, gt_norm) < 0.10

    return {"matched": matched, "gt": gt_val, "pred": pred_val}


def items_to_text(items: list) -> str:
    """แปลง items เป็นข้อความบรรทัดต่อบรรทัด สำหรับวัด CER ของตารางสินค้า"""
    return "\n".join(
        "|".join(str(it.get(c, "")) for c in schema.ITEM_FIELDS)
        for it in (items or [])
    )


def calc_items_metrics(pred_fields: dict, gt_fields: dict) -> dict:
    """
    วัดความถูกต้องของรายการสินค้า (items)
    - count_match : จำนวนแถวตรงกันไหม
    - cer         : CER ของตารางสินค้าทั้งก้อน (ทั้ง description + ตัวเลข)
    """
    gt_items   = gt_fields.get("items", []) or []
    pred_items = pred_fields.get("items", []) or []
    gt_txt     = items_to_text(gt_items)
    pred_txt   = items_to_text(pred_items)
    return {
        "count_match": len(pred_items) == len(gt_items),
        "cer":         calc_cer(pred_txt, gt_txt) if gt_txt else (0.0 if not pred_txt else 1.0),
        "gt_count":    len(gt_items),
        "pred_count":  len(pred_items),
    }


# ══════════════════════════════════════════════
# INFERENCE
# ══════════════════════════════════════════════

def load_model(model_type: str):
    """โหลด model ตาม type: 'baseline' หรือ 'finetuned'"""
    print(f"\nโหลด {model_type} model...")

    processor = ProcessorClass.from_pretrained(
        cfg.BASE_MODEL_ID,
        min_pixels=cfg.MIN_PIXELS,
        max_pixels=cfg.MAX_PIXELS,
        trust_remote_code=True,
    )

    # 4-bit quant เหมือนตอน train (Finetune.py) -- ไม่งั้น bf16 เต็ม (~6GB แค่
    # น้ำหนัก) ไม่พอ VRAM 6GB ของการ์ดนี้ HF จะ offload บาง layer ไป CPU RAM
    # อัตโนมัติ ทำให้ generate() ช้าลงมหาศาล (ส่ง tensor ข้าม GPU/CPU ทุก token)
    bnb_config = BitsAndBytesConfig(
        load_in_4bit=True,
        bnb_4bit_use_double_quant=True,
        bnb_4bit_quant_type="nf4",
        bnb_4bit_compute_dtype=torch.bfloat16,
    )

    if model_type == "baseline":
        model = ModelClass.from_pretrained(
            cfg.BASE_MODEL_ID,
            device_map={"": 0},
            quantization_config=bnb_config,
            torch_dtype=torch.bfloat16,
            trust_remote_code=True,
        )
    else:
        base = ModelClass.from_pretrained(
            cfg.BASE_MODEL_ID,
            device_map={"": 0},
            quantization_config=bnb_config,
            torch_dtype=torch.bfloat16,
            trust_remote_code=True,
        )
        model = PeftModel.from_pretrained(base, cfg.FINETUNED_DIR)

    model.eval()
    print(f"   โหลดสำเร็จ")
    return model, processor


def run_inference(model, processor, image_b64: str, system_prompt: str, user_prompt: str) -> str:
    """รัน inference กับรูปใบกำกับ 1 ใบ"""
    b64_data = image_b64.split(",")[-1]
    image = ImageOps.exif_transpose(Image.open(BytesIO(base64.b64decode(b64_data)))).convert("RGB")

    messages = [
        {"role": "system", "content": system_prompt},
        {
            "role": "user",
            "content": [
                {"type": "image",  "image": image},
                {"type": "text",   "text":  user_prompt}
            ]
        }
    ]

    text = processor.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
    inputs = processor(
        text=[text],
        images=[image],
        return_tensors="pt"
    ).to(model.device)

    with torch.no_grad():
        output_ids = model.generate(
            **inputs,
            max_new_tokens=cfg.MAX_NEW_TOKENS,
            do_sample=False,
            pad_token_id=processor.tokenizer.eos_token_id,
        )

    return processor.tokenizer.decode(
        output_ids[0][inputs["input_ids"].shape[1]:],
        skip_special_tokens=True
    )


# ══════════════════════════════════════════════
# EVALUATE
# ══════════════════════════════════════════════

def evaluate_model(model, processor, test_samples: list, model_name: str) -> list:
    """วัดผล model กับ test set ทั้งหมด"""
    print(f"\nEvaluating: {model_name}")
    results = []

    for sample in tqdm(test_samples, desc=model_name):
        messages   = sample["messages"]
        sys_prompt = messages[0]["content"]
        user_msg   = messages[1]["content"]
        gt_text    = messages[2]["content"]   # ground truth (canonical JSON string)

        image_b64  = next(c["image"] for c in user_msg if c["type"] == "image")
        user_text  = next(c["text"]  for c in user_msg if c["type"] == "text")

        # Inference
        try:
            pred_text = run_inference(model, processor, image_b64, sys_prompt, user_text)
        except Exception as e:
            pred_text = ""
            print(f"   Error: {e}")

        # parse JSON -> fields (ทนทานต่อ markdown fence / control char)
        pred_fields = schema.parse_model_json(pred_text)
        gt_fields   = sample.get("_gt_fields") or schema.parse_model_json(gt_text)

        # CER/WER วัดบน canonical JSON ของทั้งคู่
        # -> scaffolding เหมือนกันเป๊ะ ค่าต่าง = ค่าใน field ต่างจริง (ไม่ใช่ format)
        gt_canon   = schema.build_canonical_json(gt_fields)
        pred_canon = schema.build_canonical_json(pred_fields)
        sample_cer = calc_cer(pred_canon, gt_canon)
        sample_wer = calc_wer(pred_canon, gt_canon)
        emr        = 1.0 if pred_canon == gt_canon else 0.0

        # Field Accuracy (รายช่อง) + items
        field_results = {
            fld["key"]: calc_field_accuracy(pred_fields, gt_fields, fld)
            for fld in cfg.FIELDS
        }
        items_metrics = calc_items_metrics(pred_fields, gt_fields)

        results.append({
            "id":            sample.get("id", ""),
            "pred":          pred_text,
            "gt":            gt_text,
            "cer":           sample_cer,
            "wer":           sample_wer,
            "emr":           emr,
            "field_results": field_results,
            "items":         items_metrics,
            "json_ok":       bool(pred_fields and any(pred_fields.get(k) for k in schema.SCALAR_FIELDS)),
        })

    return results


def summarize(results: list, model_name: str) -> dict:
    """สรุปผลเฉลี่ยทั้ง test set"""
    n = len(results)
    if n == 0:
        return {}

    avg_cer = sum(r["cer"] for r in results) / n
    avg_wer = sum(r["wer"] for r in results) / n
    avg_emr = sum(r["emr"] for r in results) / n

    # Field accuracy แยกแต่ละ field
    field_acc = {}
    for fld in cfg.FIELDS:
        key = fld["key"]
        matched = [r["field_results"][key]["matched"]
                   for r in results
                   if r["field_results"][key]["matched"] is not None]
        field_acc[key] = sum(matched) / len(matched) if matched else None

    avg_field = sum(v for v in field_acc.values() if v is not None)
    avg_field /= sum(1 for v in field_acc.values() if v is not None) or 1

    # Items + JSON validity
    items_cer        = sum(r["items"]["cer"] for r in results) / n
    items_count_acc  = sum(1 for r in results if r["items"]["count_match"]) / n
    json_valid_rate  = sum(1 for r in results if r["json_ok"]) / n

    return {
        "model":           model_name,
        "n_samples":       n,
        "avg_cer":         avg_cer,
        "avg_wer":         avg_wer,
        "avg_emr":         avg_emr,
        "avg_field":       avg_field,
        "field_acc":       field_acc,
        "items_cer":       items_cer,
        "items_count_acc": items_count_acc,
        "json_valid_rate": json_valid_rate,
    }


# ══════════════════════════════════════════════
# EXPORT EXCEL REPORT
# ══════════════════════════════════════════════

def export_excel(baseline_sum: dict, ft_sum: dict, baseline_res: list, ft_res: list):
    """Export รายงานเปรียบเทียบเป็น Excel"""
    Path(cfg.OUTPUT_DIR).mkdir(exist_ok=True)
    ts   = datetime.now().strftime("%Y%m%d_%H%M")
    path = f"{cfg.OUTPUT_DIR}/ocr_evaluation_{ts}.xlsx"

    wb = openpyxl.Workbook()

    # ── Sheet 1: สรุป ──
    ws1 = wb.active
    ws1.title = "สรุปผล"

    # สี
    blue_fill   = PatternFill("solid", fgColor="1F4E79")
    green_fill  = PatternFill("solid", fgColor="1D6A3A")
    amber_fill  = PatternFill("solid", fgColor="7C4700")
    gray_fill   = PatternFill("solid", fgColor="2D2D2D")
    hdr_font    = Font(color="FFFFFF", bold=True, size=11)
    num_fmt     = "0.00%"

    def hdr(ws, row, col, val, fill):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = fill
        c.font = hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        return c

    def val(ws, row, col, v, fmt=None, bold=False, color=None):
        c = ws.cell(row=row, column=col, value=v)
        if fmt:   c.number_format = fmt
        if bold:  c.font = Font(bold=True, color=color or "000000")
        c.alignment = Alignment(horizontal="center")
        return c

    # Title
    ws1.merge_cells("A1:F1")
    t = ws1["A1"]
    t.value = "รายงานเปรียบเทียบ: Baseline vs Fine-tuned Typhoon OCR"
    t.fill  = PatternFill("solid", fgColor="0D1B2A")
    t.font  = Font(color="FFFFFF", bold=True, size=14)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 32

    # Headers row
    headers = ["Metric", "Baseline", "Fine-tuned", "ผลต่าง", "ดีขึ้น?", "หมายเหตุ"]
    for i, h in enumerate(headers, 1):
        hdr(ws1, 2, i, h, blue_fill)
    ws1.row_dimensions[2].height = 24

    # ── Overall Metrics ──
    rows = [
        ("CER (↓ ดีกว่า)",   baseline_sum["avg_cer"], ft_sum["avg_cer"],  False, "ยิ่งน้อยยิ่งดี"),
        ("WER (↓ ดีกว่า)",   baseline_sum["avg_wer"], ft_sum["avg_wer"],  False, "ยิ่งน้อยยิ่งดี"),
        ("EMR (↑ ดีกว่า)",   baseline_sum["avg_emr"], ft_sum["avg_emr"],  True,  "Exact Match Rate"),
        ("Field Acc (↑)",     baseline_sum["avg_field"],ft_sum["avg_field"],True, "เฉลี่ยทุก field"),
    ]

    for i, (name, base_v, ft_v, higher_better, note) in enumerate(rows, 3):
        diff   = ft_v - base_v
        better = diff > 0 if higher_better else diff < 0
        icon   = "" if better else ("" if abs(diff) < 0.005 else "")

        ws1.cell(row=i, column=1, value=name)
        val(ws1, i, 2, base_v, num_fmt)
        val(ws1, i, 3, ft_v,   num_fmt)
        val(ws1, i, 4, diff,   num_fmt, bold=True, color="1D6A3A" if better else "7C0A02")
        val(ws1, i, 5, icon)
        val(ws1, i, 6, note)

    # ── Field Accuracy ──
    ws1.cell(row=8, column=1, value="── Field Accuracy ──").font = Font(bold=True, color="1F4E79")

    for i, fld in enumerate(cfg.FIELDS, 9):
        key  = fld["key"]
        base_fa = baseline_sum["field_acc"].get(key)
        ft_fa   = ft_sum["field_acc"].get(key)
        if base_fa is None or ft_fa is None:
            continue

        diff   = ft_fa - base_fa
        better = diff > 0
        icon   = "" if better else ("" if abs(diff) < 0.01 else "")

        ws1.cell(row=i, column=1, value=fld["label"])
        val(ws1, i, 2, base_fa, num_fmt)
        val(ws1, i, 3, ft_fa,   num_fmt)
        val(ws1, i, 4, diff,    num_fmt, bold=True, color="1D6A3A" if better else "7C0A02")
        val(ws1, i, 5, icon)

    # column widths
    for col, width in zip("ABCDEF", [22, 12, 12, 12, 10, 22]):
        ws1.column_dimensions[chr(64+["A","B","C","D","E","F"].index(col)+1)].width = width

    # ── Sheet 2: รายละเอียดรายใบ ──
    ws2 = wb.create_sheet("รายละเอียดรายใบ")
    detail_headers = ["ID", "CER Baseline", "CER Fine-tuned", "WER Baseline", "WER Fine-tuned",
                      "Field Acc Baseline", "Field Acc Fine-tuned", "ดีขึ้น?"]
    for i, h in enumerate(detail_headers, 1):
        hdr(ws2, 1, i, h, blue_fill)

    for r_idx, (br, fr) in enumerate(zip(baseline_res, ft_res), 2):
        base_fa = sum(
            1 for fld in cfg.FIELDS
            if br["field_results"][fld["key"]]["matched"] is True
        ) / max(sum(
            1 for fld in cfg.FIELDS
            if br["field_results"][fld["key"]]["matched"] is not None
        ), 1)
        ft_fa = sum(
            1 for fld in cfg.FIELDS
            if fr["field_results"][fld["key"]]["matched"] is True
        ) / max(sum(
            1 for fld in cfg.FIELDS
            if fr["field_results"][fld["key"]]["matched"] is not None
        ), 1)

        better = (fr["cer"] < br["cer"]) and (ft_fa >= base_fa)
        icon   = "" if better else ("" if abs(fr["cer"] - br["cer"]) < 0.01 else "")

        row_data = [br["id"], br["cer"], fr["cer"], br["wer"], fr["wer"], base_fa, ft_fa, icon]
        for c_idx, v in enumerate(row_data, 1):
            c = ws2.cell(row=r_idx, column=c_idx, value=v)
            if c_idx in [2, 3, 4, 5, 6, 7]:
                c.number_format = num_fmt
            c.alignment = Alignment(horizontal="center")

    # ── Sheet 3: ข้อความดิบ (GT vs Baseline vs Fine-tuned) ──
    ws3 = wb.create_sheet("ข้อความดิบ")
    raw_headers = ["ID", "Ground Truth", "Baseline Pred", "Fine-tuned Pred"]
    for i, h in enumerate(raw_headers, 1):
        hdr(ws3, 1, i, h, blue_fill)

    for r_idx, (br, fr) in enumerate(zip(baseline_res, ft_res), 2):
        row_data = [br["id"], br["gt"], br["pred"], fr["pred"]]
        for c_idx, v in enumerate(row_data, 1):
            c = ws3.cell(row=r_idx, column=c_idx, value=v)
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for col, width in zip("ABCD", [12, 60, 60, 60]):
        ws3.column_dimensions[col].width = width

    wb.save(path)
    print(f"\nExport Excel → {path}")
    return path


# ══════════════════════════════════════════════
# PRINT SUMMARY TABLE
# ══════════════════════════════════════════════

def print_summary(baseline_sum: dict, ft_sum: dict):
    """แสดงผลสรุปใน terminal"""
    print(f"\n{'='*60}")
    print(f"ผลการประเมิน — Baseline vs Fine-tuned")
    print(f"{'='*60}")
    print(f"{'Metric':<20} {'Baseline':>12} {'Fine-tuned':>12} {'ผลต่าง':>10}")
    print(f"{'─'*60}")

    metrics = [
        ("CER (↓)",         "avg_cer",         False),
        ("WER (↓)",         "avg_wer",         False),
        ("EMR (↑)",         "avg_emr",         True),
        ("Field Acc (↑)",   "avg_field",       True),
        ("Items CER (↓)",   "items_cer",       False),
        ("Items count (↑)", "items_count_acc", True),
        ("JSON valid (↑)",  "json_valid_rate", True),
    ]

    for label, key, higher_better in metrics:
        b = baseline_sum[key]
        f = ft_sum[key]
        d = f - b
        better = d > 0 if higher_better else d < 0
        icon = "" if better else ("" if abs(d) < 0.005 else "")
        print(f"{label:<20} {b:>11.1%} {f:>11.1%} {d:>+10.1%} {icon}")

    print(f"\n{'─'*60}")
    print(f"Field-level Accuracy:")
    print(f"{'─'*60}")
    print(f"{'Field':<22} {'Baseline':>10} {'Fine-tuned':>10} {'ผลต่าง':>8}")

    for fld in cfg.FIELDS:
        key  = fld["key"]
        b    = baseline_sum["field_acc"].get(key)
        f    = ft_sum["field_acc"].get(key)
        if b is None or f is None:
            continue
        d    = f - b
        icon = "" if d > 0.01 else ("" if abs(d) <= 0.01 else "")
        print(f"  {fld['label']:<20} {b:>9.1%} {f:>9.1%} {d:>+7.1%} {icon}")

    print(f"{'='*60}")


# ══════════════════════════════════════════════
# MAIN
# ══════════════════════════════════════════════

def main():
    # ตรวจสอบ
    if not Path(cfg.TEST_JSONL).exists():
        print(f"ไม่พบ {cfg.TEST_JSONL} — รัน format_dataset.py ก่อน")
        return

    if not Path(cfg.FINETUNED_DIR).exists():
        print(f"ไม่พบ {cfg.FINETUNED_DIR} — รัน finetune.py ก่อน")
        return

    if not torch.cuda.is_available():
        print("ไม่พบ GPU")
        return

    # โหลด test set
    print("โหลด Test Set...")
    test_samples = []
    with open(cfg.TEST_JSONL, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                test_samples.append(json.loads(line))

    print(f"   {len(test_samples)} samples")

    Path(cfg.OUTPUT_DIR).mkdir(exist_ok=True)

    # ── Evaluate Baseline ──
    base_model, base_proc = load_model(model_type="baseline")
    baseline_results = evaluate_model(base_model, base_proc, test_samples, "Baseline")
    baseline_summary = summarize(baseline_results, "Baseline")

    # คืน VRAM ก่อนโหลดโมเดลถัดไป
    del base_model
    torch.cuda.empty_cache()

    # ── Evaluate Fine-tuned ──
    ft_model, ft_proc = load_model(model_type="finetuned")
    ft_results = evaluate_model(ft_model, ft_proc, test_samples, "Fine-tuned")
    ft_summary = summarize(ft_results, "Fine-tuned")

    del ft_model
    torch.cuda.empty_cache()

    # ── แสดงผล ──
    print_summary(baseline_summary, ft_summary)

    # ── Export ──
    excel_path = export_excel(baseline_summary, ft_summary, baseline_results, ft_results)

    # บันทึก JSON ด้วย
    json_path = f"{cfg.OUTPUT_DIR}/results_{datetime.now().strftime('%Y%m%d_%H%M')}.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump({
            "baseline": baseline_summary,
            "finetuned": ft_summary,
            "reference": {
                "CER_benchmark": "CER < 2% = ดีมาก, < 5% = ดี, < 10% = พอใช้",
                "source": "Docsumo OCR Accuracy Guide 2025, arXiv:2601.14722"
            }
        }, f, ensure_ascii=False, indent=2)

    print(f"JSON results → {json_path}")
    print(f"\nEvaluation เสร็จแล้ว!")
    print(f"   ใช้ผลลัพธ์จาก Excel ใส่ใน Thesis ได้เลยครับ")


if __name__ == "__main__":
    main()